#!/usr/bin/env python3
"""
ETL Pipeline: Spreadsheet -> Normalized JSON database
=====================================================
The spreadsheet IS the CMS. Rerun this script whenever the sheet changes:
    python3 etl/build_data.py "<path-to-xlsx>"
It regenerates src/data/laptops.json which the site consumes at build time.

Pipeline stages:
 1. Parse multi-row header Excel (sections = price brackets)
 2. Extract embedded Amazon hyperlinks (affiliate tag preserved)
 3. Normalize brands / GPU names / panel types (fix typos like 'NIVIDIA')
 4. Derive structured fields (resolution class, storage GB, PPI, etc.)
 5. Compute Smart Ranking scores (overall, gaming, office, programming,
    engineering, creator, student, travel, ai, budget/value)
 6. Assign badges (Editor's Choice, Best Value, Highly Rated, ...)
 7. Emit normalized JSON
"""
import sys, json, re, math, unicodedata
import openpyxl

SRC = sys.argv[1] if len(sys.argv) > 1 else "/home/user/uploaded_files/Laptops Recs August 2026 (2).xlsx"
OUT = "/home/user/webapp/src/data/laptops.json"

wb = openpyxl.load_workbook(SRC)
ws = wb["Main"]

def cell(r, c):
    v = ws.cell(row=r, column=c).value
    if isinstance(v, str):
        v = v.strip()
        if v in ("-", "", "N/A", "n/a"):
            return None
    return v

def slugify(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+", "-", s).strip("-").lower()
    return re.sub(r"-{2,}", "-", s)

BRAND_FIX = {"asus": "ASUS", "hp": "HP", "msi": "MSI", "lg": "LG"}
def norm_brand(b):
    if not b: return None
    b = b.strip()
    return BRAND_FIX.get(b.lower(), b.title() if b.islower() or b.isupper() and len(b) > 3 else b)

def norm_gpu_brand(b):
    if not b: return None
    b = b.strip()
    if b.upper() in ("NVIDIA", "NIVIDIA"): return "NVIDIA"
    return b

def parse_storage_gb(v):
    if v is None: return None
    s = str(v).upper().replace(",", "")
    m = re.search(r"([\d.]+)\s*TB", s)
    if m: return int(float(m.group(1)) * 1024)
    m = re.search(r"([\d.]+)\s*GB", s)
    if m: return int(float(m.group(1)))
    try: return int(float(s))
    except: return None

rows = []
current_bracket = None
r = 2  # header rows: 1 (groups), 2 (fields); data starts row 3-ish
for r in range(3, ws.max_row + 1):
    brand = cell(r, 1)
    price = cell(r, 38)
    nonnull = sum(1 for c in range(1, 39) if cell(r, c) is not None)
    if brand and price is None and nonnull <= 2:
        if str(brand).lower().startswith("under"):
            current_bracket = str(brand)
        continue
    if not brand or price is None:
        continue
    link_cell = ws.cell(row=r, column=35)
    amazon = link_cell.hyperlink.target if link_cell.hyperlink else None
    rows.append({
        "row": r, "bracket": current_bracket,
        "brand": norm_brand(brand), "model": str(cell(r, 2) or "").strip(),
        "segment": cell(r, 3), "type": cell(r, 4),
        "screen_size": cell(r, 6), "resolution": cell(r, 7),
        "refresh": cell(r, 8), "panel": (cell(r, 9) or "").strip() or None,
        "touch": cell(r, 10),
        "cpu_brand": (cell(r, 12) or "").strip() or None,
        "cpu_model": (cell(r, 13) or "").strip() or None,
        "cpu_cores": cell(r, 14), "cpu_mt": cell(r, 15),
        "passmark": cell(r, 16),
        "gpu_kind": cell(r, 18),
        "gpu_brand": norm_gpu_brand(cell(r, 19)),
        "gpu_model": (str(cell(r, 20)).strip() if cell(r, 20) else None),
        "g3dmark": cell(r, 21),
        "ram_gb": cell(r, 23), "ram_type": (cell(r, 24) or "").strip() or None,
        "ram_soldered": cell(r, 25),
        "storage_raw": cell(r, 27), "storage_type": cell(r, 28),
        "depth_in": cell(r, 30), "length_in": cell(r, 31),
        "width_in": cell(r, 32), "weight_lbs": cell(r, 33),
        "amazon": amazon, "rating": cell(r, 36),
        "reviews": cell(r, 37), "price": price,
    })

print(f"Parsed {len(rows)} laptops")

# ---------- Normalize & derive ----------
def yes(v): return str(v).strip().lower() == "yes" if v is not None else False

def res_class(res, panel):
    if not res: return None
    try:
        w, h = [int(x) for x in str(res).lower().split("x")]
    except: return None
    px = w
    if px >= 3800: return "4K"
    if px >= 3000: return "3K+"
    if px >= 2500: return "QHD+"
    if px >= 1900:
        return "FHD+" if h > 1100 else "FHD"
    return "HD"

laptops = []
slug_seen = {}
for i, x in enumerate(rows):
    name = f"{x['brand']} {x['model']}".strip()
    base = slugify(name)
    n = slug_seen.get(base, 0)
    slug_seen[base] = n + 1
    slug = base if n == 0 else f"{base}-{n+1}"
    storage_gb = parse_storage_gb(x["storage_raw"])
    try:
        w, h = [int(v) for v in str(x["resolution"]).lower().split("x")]
        ppi = round(math.sqrt(w*w + h*h) / float(x["screen_size"])) if x["screen_size"] else None
    except:
        w = h = ppi = None
    lp = {
        "id": i + 1, "slug": slug, "name": name,
        "brand": x["brand"], "model": x["model"],
        "segment": x["segment"] or "General",
        "formFactor": x["type"] or "Traditional",
        "priceBracket": x["bracket"],
        "price": round(float(x["price"]), 2),
        "display": {
            "sizeInches": x["screen_size"],
            "resolution": x["resolution"],
            "resW": w, "resH": h, "ppi": ppi,
            "resClass": res_class(x["resolution"], x["panel"]),
            "refreshHz": x["refresh"],
            "panel": x["panel"],
            "touch": yes(x["touch"]),
        },
        "cpu": {
            "brand": x["cpu_brand"], "model": x["cpu_model"],
            "cores": x["cpu_cores"], "multiThread": yes(x["cpu_mt"]),
            "passmark": x["passmark"],
        },
        "gpu": {
            "kind": x["gpu_kind"] or "Integrated",
            "dedicated": (x["gpu_kind"] == "Dedicated"),
            "brand": x["gpu_brand"], "model": x["gpu_model"],
            "g3dmark": x["g3dmark"] if isinstance(x["g3dmark"], (int, float)) else None,
        },
        "ram": {"gb": x["ram_gb"], "type": x["ram_type"], "soldered": yes(x["ram_soldered"])},
        "storage": {"gb": storage_gb, "type": x["storage_type"], "raw": str(x["storage_raw"])},
        "physical": {
            "thicknessIn": x["depth_in"], "lengthIn": x["length_in"],
            "widthIn": x["width_in"], "weightLbs": x["weight_lbs"],
            "weightKg": round(x["weight_lbs"] * 0.4536, 2) if isinstance(x["weight_lbs"], (int, float)) else None,
        },
        "amazon": {"url": x["amazon"], "rating": x["rating"], "reviewCount": x["reviews"]},
    }
    laptops.append(lp)

# ---------- Smart Ranking Engine ----------
def collect(vals):
    vals = [v for v in vals if isinstance(v, (int, float))]
    return (min(vals), max(vals)) if vals else (0, 1)

pm_min, pm_max = collect([l["cpu"]["passmark"] for l in laptops])
g3_vals = [l["gpu"]["g3dmark"] for l in laptops if l["gpu"]["g3dmark"]]
g3_min, g3_max = collect(g3_vals)
pr_min, pr_max = collect([l["price"] for l in laptops])

def nlog(v, lo, hi):
    """log-scaled 0..1 normalization (benchmarks follow log perception)"""
    if not isinstance(v, (int, float)) or v <= 0: return 0.0
    lo = max(lo, 1)
    return max(0.0, min(1.0, (math.log(v) - math.log(lo)) / (math.log(hi) - math.log(lo))))

def nlin(v, lo, hi, invert=False):
    if not isinstance(v, (int, float)): return 0.0
    t = max(0.0, min(1.0, (v - lo) / (hi - lo)))
    return 1 - t if invert else t

for l in laptops:
    cpu = nlog(l["cpu"]["passmark"], pm_min, pm_max)
    gpu = nlog(l["gpu"]["g3dmark"], g3_min, g3_max) if l["gpu"]["g3dmark"] else (0.18 * cpu + 0.10)
    if not l["gpu"]["dedicated"]:
        gpu = min(gpu, 0.42)  # iGPU ceiling
    ram = nlin(l["ram"]["gb"], 4, 64)
    sto = nlin(l["storage"]["gb"] or 0, 64, 4096)
    d = l["display"]
    scr = (0.30 * nlin(d["ppi"] or 100, 100, 260)
           + 0.28 * nlin(d["refreshHz"] or 60, 60, 300)
           + 0.28 * {"OLED": 1.0, "AMOLED": 1.0, "Mini LED": 0.92, "Liquid Retina": 0.85, "IPS": 0.6}.get(d["panel"], 0.4)
           + 0.14 * (1.0 if d["touch"] else 0.6))
    wt = nlin(l["physical"]["weightLbs"], 2.0, 8.0, invert=True)
    thin = nlin(l["physical"]["thicknessIn"] or 1.0, 0.4, 1.2, invert=True)
    rat = nlin(l["amazon"]["rating"] or 3.5, 3.0, 5.0)
    rev_conf = nlin(math.log10(max(l["amazon"]["reviewCount"] or 1, 1)), 0, 4)
    rating_w = rat * (0.5 + 0.5 * rev_conf)
    price_pos = nlog(l["price"], pr_min, pr_max)   # 0 cheap .. 1 expensive
    ram_up = 0.0 if l["ram"]["soldered"] else 1.0

    perf = 0.55 * cpu + 0.45 * gpu
    value = max(0.0, min(1.0, perf + 0.25 * ram + 0.15 * sto - 0.95 * price_pos + 0.42))

    def pct(x): return round(max(1.0, min(10.0, x * 10)), 1)
    scores = {
        "overall":     pct(0.26*cpu + 0.20*gpu + 0.12*ram + 0.08*sto + 0.14*scr + 0.06*wt + 0.14*rating_w),
        "gaming":      pct(0.46*gpu + 0.22*cpu + 0.10*ram + 0.14*nlin(d["refreshHz"] or 60, 60, 300) + 0.08*sto),
        "office":      pct(0.30*cpu + 0.16*ram + 0.10*sto + 0.14*scr + 0.16*wt + 0.14*rating_w),
        "programming": pct(0.34*cpu + 0.24*ram + 0.14*sto + 0.14*scr + 0.08*wt + 0.06*(1 if (d["resH"] or 0) >= 1200 else 0)),
        "engineering": pct(0.30*cpu + 0.30*gpu + 0.20*ram + 0.12*sto + 0.08*scr),
        "creator":     pct(0.26*cpu + 0.26*gpu + 0.14*ram + 0.10*sto + 0.24*scr),
        "student":     pct(0.20*cpu + 0.14*ram + 0.20*wt + 0.12*scr + 0.20*(1-price_pos) + 0.14*rating_w),
        "travel":      pct(0.34*wt + 0.20*thin + 0.14*cpu + 0.12*scr + 0.10*ram + 0.10*rating_w),
        "ai":          pct(0.34*gpu + 0.28*cpu + 0.26*ram + 0.12*sto),
        "value":       pct(value),
    }
    l["scores"] = scores
    l["_internals"] = {"cpuN": round(cpu,3), "gpuN": round(gpu,3), "ramN": round(ram,3),
                       "stoN": round(sto,3), "scrN": round(scr,3), "wtN": round(wt,3),
                       "ratN": round(rating_w,3), "priceN": round(price_pos,3), "ramUp": ram_up}

# ---------- Badges ----------
def top_ids(key, n=5, flt=None):
    pool = [l for l in laptops if (flt(l) if flt else True)]
    pool.sort(key=lambda l: -l["scores"][key])
    return [l["id"] for l in pool[:n]]

badges = {}
def add_badge(lid, b):
    badges.setdefault(lid, []).append(b)

for lid in top_ids("overall", 3): add_badge(lid, "Editor's Choice")
for lid in top_ids("value", 4): add_badge(lid, "Best Value")
for lid in top_ids("gaming", 3, lambda l: l["gpu"]["dedicated"]): add_badge(lid, "Top Gaming Pick")
hr = [l["id"] for l in laptops if (l["amazon"]["rating"] or 0) >= 4.4 and (l["amazon"]["reviewCount"] or 0) >= 500]
for lid in hr: add_badge(lid, "Highly Rated")
pop = sorted(laptops, key=lambda l: -(l["amazon"]["reviewCount"] or 0))[:5]
for l in pop: add_badge(l["id"], "Popular")
cheap = min(laptops, key=lambda l: l["price"])
add_badge(cheap["id"], "Lowest Price")
for l in laptops:
    l["badges"] = list(dict.fromkeys(badges.get(l["id"], [])))

# ---------- Emit ----------
import os
os.makedirs(os.path.dirname(OUT), exist_ok=True)
meta = {
    "generatedFrom": SRC.split("/")[-1],
    "updated": "2026-08-19",
    "count": len(laptops),
    "affiliateDisclosure": "As Amazon Associates we may earn commission from qualifying purchases.",
}
with open(OUT, "w") as f:
    json.dump({"meta": meta, "laptops": laptops}, f, ensure_ascii=False)
print(f"Wrote {OUT}: {len(laptops)} laptops")
print("Missing amazon links:", sum(1 for l in laptops if not l["amazon"]["url"]))
print("Sample slugs:", [l["slug"] for l in laptops[:5]])
